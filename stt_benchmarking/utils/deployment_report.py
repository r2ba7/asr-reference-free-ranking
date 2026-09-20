#!/usr/bin/env python3
"""Combine every CSV (and JSON scalar file) written by DeploymentProtocol into one DOCX and one PDF, each file as a table.
Tables are ordered by what they answer (ranking first, raw enumeration last) rather than alphabetically.
Usage: python deployment_report.py <directory> [-n NAME] [-s SUFFIX] [--decimals 4] [--max-rows 60] [--recursive] [--no-docx] [--no-pdf] [--portrait]
Outputs are always written into <directory> as <NAME><SUFFIX>.docx / .pdf
Deps:  pip install pandas python-docx reportlab
"""
import argparse,json,math,re,sys
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4,landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate,Table,TableStyle,Paragraph,Spacer

# reading order: the reported ranking, then its confidence, then what the weights
# were built from, then the raw per-subset enumeration everything is derived from
ORDER=["scalars","ranking","model_confidence","rank_distribution","pair_probability",
       "model_filtration","by_pool_size","cross_wer","diagnostics_runs",
       "diagnostics_dropout","subpool_filtration","subpool_ranks"]

def okey(p):
    """Order by ORDER on the stem's leading table name, unknown files last, then natural sort."""
    s=p.stem.lower()
    hit=[i for i,n in enumerate(ORDER) if s.startswith(n)]
    rank=min(hit) if hit else len(ORDER)
    return [rank]+[int(t) if t.isdigit() else t for t in re.split(r"(\d+)",s)]

def fmt(v,nd):
    if v is None:return ""
    if isinstance(v,float):
        if math.isnan(v):return ""
        if math.isinf(v):return "inf" if v>0 else "-inf"
        if v!=0 and abs(v)<10**(-nd):return f"{v:.{nd}e}"
        return f"{v:.{nd}f}".rstrip("0").rstrip(".") or "0"
    if isinstance(v,bool):return str(v)
    return str(v)

def flatten(o,prefix=""):
    out={}
    if isinstance(o,dict):
        if not o:return {prefix or "value":"{}"}
        for k,v in o.items():out.update(flatten(v,f"{prefix}.{k}" if prefix else str(k)))
    elif isinstance(o,list):
        if not o:return {prefix or "value":"[]"}
        for i,v in enumerate(o):out.update(flatten(v,f"{prefix}[{i}]"))
    else:out[prefix or "value"]=o
    return out

def mapdf(df,fn):
    df=df.astype(object)
    return df.map(fn) if hasattr(df,"map") else df.applymap(fn)

def load(path,nd):
    if path.suffix.lower()==".json":
        data=json.loads(path.read_text(encoding="utf-8"))
        flat=flatten(data)
        return pd.DataFrame({"key":list(flat.keys()),"value":[fmt(v,nd) for v in flat.values()]})
    df=pd.read_csv(path)
    df=df.rename(columns={c:"" for c in df.columns if str(c).startswith("Unnamed:")})
    df=df.where(pd.notnull(df),None)
    return mapdf(df,lambda v:fmt(v,nd))

def collect(root,recursive):
    pat="**/*" if recursive else "*"
    files=[p for p in root.glob(pat) if p.suffix.lower() in(".csv",".json") and p.is_file()]
    return sorted(files,key=okey)

def build_docx(tables,out,title,portrait):
    doc=Document()
    sec=doc.sections[0]
    if not portrait:
        sec.orientation=WD_ORIENT.LANDSCAPE
        sec.page_width,sec.page_height=sec.page_height,sec.page_width
    for m in("left_margin","right_margin","top_margin","bottom_margin"):setattr(sec,m,Inches(0.5))
    style=doc.styles["Normal"]
    style.font.name="Calibri"
    style.font.size=Pt(9)
    h=doc.add_heading(title,level=0)
    h.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for i,(name,df) in enumerate(tables):
        doc.add_heading(name,level=1)
        t=doc.add_table(rows=1,cols=len(df.columns))
        t.style="Table Grid"
        t.autofit=True
        fs=Pt(7 if len(df.columns)>8 else 8)
        for j,c in enumerate(df.columns):
            cell=t.rows[0].cells[j]
            cell.text=""
            r=cell.paragraphs[0].add_run(str(c))
            r.bold=True
            r.font.size=fs
            r.font.color.rgb=RGBColor(0,0,0)
        for _,row in df.iterrows():
            cells=t.add_row().cells
            for j,v in enumerate(row):
                cells[j].text=""
                r=cells[j].paragraphs[0].add_run(str(v))
                r.font.size=fs
        if i<len(tables)-1:doc.add_page_break()
    doc.save(out)

def col_widths(df,avail):
    w=[]
    for c in df.columns:
        m=len(str(c))
        if len(df):m=max(m,int(df[c].astype(str).str.len().max()))
        w.append(min(max(m,4),45))
    tot=sum(w)
    return [max(avail*x/tot,14*mm) for x in w]

def build_pdf(tables,out,title,portrait):
    page=A4 if portrait else landscape(A4)
    doc=SimpleDocTemplate(str(out),pagesize=page,leftMargin=12*mm,rightMargin=12*mm,topMargin=12*mm,bottomMargin=12*mm,title=title)
    avail=page[0]-24*mm
    h1=ParagraphStyle("h1",fontName="Helvetica-Bold",fontSize=16,spaceAfter=10)
    h2=ParagraphStyle("h2",fontName="Helvetica-Bold",fontSize=11,spaceBefore=8,spaceAfter=6)
    story=[Paragraph(title,h1),Spacer(1,4*mm)]
    for name,df in tables:
        fs=6 if len(df.columns)>10 else 7
        cs=ParagraphStyle("c",fontName="Helvetica",fontSize=fs,leading=fs+1.5)
        hs=ParagraphStyle("hh",fontName="Helvetica-Bold",fontSize=fs,leading=fs+1.5,textColor=colors.white)
        data=[[Paragraph(str(c),hs) for c in df.columns]]
        data+=[[Paragraph(str(v),cs) for v in row] for row in df.itertuples(index=False)]
        widths=col_widths(df,avail)
        scale=avail/sum(widths)
        widths=[w*scale for w in widths]
        t=Table(data,colWidths=widths,repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#37474F")),
            ("GRID",(0,0),(-1,-1),0.4,colors.HexColor("#B0BEC5")),
            ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
            ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F2F5F7")]),
            ("LEFTPADDING",(0,0),(-1,-1),3),("RIGHTPADDING",(0,0),(-1,-1),3),
            ("TOPPADDING",(0,0),(-1,-1),2),("BOTTOMPADDING",(0,0),(-1,-1),2),
        ]))
        story+=[Paragraph(name,h2),t,Spacer(1,6*mm)]
    doc.build(story)

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("directory")
    ap.add_argument("-n","--name",default=None,help="output file name, no extension (default: input directory name)")
    ap.add_argument("-s","--suffix",default="",help="string appended to name, e.g. _report or _v2")
    ap.add_argument("--max-rows",type=int,default=60,help="truncate tables to this many rows (0 = no limit); subpool_ranks and subpool_filtration run to thousands of rows")
    ap.add_argument("--decimals",type=int,default=4)
    ap.add_argument("--recursive",action="store_true")
    ap.add_argument("--portrait",action="store_true")
    ap.add_argument("--no-docx",action="store_true")
    ap.add_argument("--no-pdf",action="store_true")
    ap.add_argument("--title",default=None)
    a=ap.parse_args()
    root=Path(a.directory).expanduser().resolve()
    if not root.is_dir():sys.exit(f"not a directory: {root}")
    files=collect(root,a.recursive)
    if not files:sys.exit(f"no csv/json files in {root}")
    tables=[];skipped=[]
    for p in files:
        try:
            df=load(p,a.decimals)
            if df.empty:
                skipped.append((p.name,"empty after parsing"));continue
            n=len(df)
            name=p.stem
            if a.max_rows and n>a.max_rows:
                df=df.head(a.max_rows)
                name=f"{p.stem}  (first {a.max_rows} of {n} rows)"
            tables.append((name,df))
            print(f"loaded {p.name}: {n}x{df.shape[1]}"+(f" -> truncated to {a.max_rows}" if a.max_rows and n>a.max_rows else ""))
        except Exception as e:
            skipped.append((p.name,f"{type(e).__name__}: {e}"))
    if skipped:
        print("\n"+"!"*60,file=sys.stderr)
        for n,r in skipped:print(f"SKIPPED {n} -> {r}",file=sys.stderr)
        print("!"*60+"\n",file=sys.stderr)
    print(f"{len(tables)} of {len(files)} files included")
    if not tables:sys.exit("nothing loadable")
    stem=(a.name or root.name)+a.suffix
    stem=re.sub(r'[<>:"/\\|?*]',"_",stem).strip() or root.name
    title=a.title or stem
    if not a.no_docx:
        d=root/f"{stem}.docx";build_docx(tables,d,title,a.portrait);print(f"wrote {d}")
    if not a.no_pdf:
        f=root/f"{stem}.pdf";build_pdf(tables,f,title,a.portrait);print(f"wrote {f}")

if __name__=="__main__":
    main()
